"""Word export of a print document.

The browser lays the document out and sends its print view as a
:class:`WordExport`: paragraphs of runs, and tables placed on a grid of
columns. All the lengths are in points; the widths of the columns only count
relative to each other, the tables filling the width they are given.
"""

import base64
import binascii
from io import BytesIO
from typing import Annotated, Literal

import msgspec
from docx import Document as new_document
from docx.document import Document
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import (
    WD_ALIGN_PARAGRAPH,
    WD_TAB_ALIGNMENT,
    WD_TAB_LEADER,
    WD_UNDERLINE,
)
from docx.exceptions import InvalidSpanError
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.oxml.xmlchemy import BaseOxmlElement
from docx.shared import Emu, Length as WordLength, Mm, Pt, RGBColor
from docx.table import Table as WordTable, _Cell
from docx.text.paragraph import Paragraph as WordParagraph
from docx.text.run import Run as WordRun
from PIL import Image, UnidentifiedImageError

from data.print_documents.fonts import PRINT_FONT_FAMILY, PRINT_FONT_FILES
from data.print_documents.page_margins import PageMargins
from data.print_documents.word_font_embedding import embed_font

Length = Annotated[float, msgspec.Meta(ge=0, le=10_000)]
FontSize = Annotated[float, msgspec.Meta(gt=0, le=200)]
Color = Annotated[str, msgspec.Meta(pattern='^[0-9A-Fa-f]{6}$')]
Count = Annotated[int, msgspec.Meta(ge=1, le=1_000)]
Index = Annotated[int, msgspec.Meta(ge=0, le=1_000)]
FontName = Annotated[str, msgspec.Meta(min_length=1, max_length=100)]


class Border(msgspec.Struct):
    width: Length
    style: Literal['single', 'dotted', 'dashed', 'double'] = 'single'
    color: Color = '000000'


class Borders(msgspec.Struct):
    top: Border | None = None
    right: Border | None = None
    bottom: Border | None = None
    left: Border | None = None


class TextRun(msgspec.Struct, tag_field='type', tag='text'):
    text: str
    bold: bool = False
    italic: bool = False
    underline: bool = False
    strike: bool = False
    size: FontSize | None = None
    color: Color | None = None
    font: FontName | None = None


class BlankRun(msgspec.Struct, tag_field='type', tag='blank'):
    """An empty field, drawn as a dotted line to write on; without a width, the
    line runs to the end of the paragraph."""

    width: Length | None = None
    size: FontSize | None = None
    font: FontName | None = None


class BreakRun(msgspec.Struct, tag_field='type', tag='break'):
    pass


class ImageRun(msgspec.Struct, tag_field='type', tag='image'):
    src: str
    width: Length
    height: Length


Run = TextRun | BlankRun | BreakRun | ImageRun


class Paragraph(msgspec.Struct, tag_field='type', tag='paragraph'):
    runs: list[Run]
    align: Literal['left', 'center', 'right', 'justify'] = 'left'
    space_before: Length = 0


class Cell(msgspec.Struct):
    row: Index
    column: Index
    blocks: list['Block']
    column_span: Count = 1
    row_span: Count = 1
    borders: Borders | None = None
    fill: Color | None = None
    vertical_align: Literal['top', 'center', 'bottom'] = 'top'
    space_before: Length = 0
    padding: Length = 0
    """The space between the content of the cell and its top and bottom."""


class Table(msgspec.Struct, tag_field='type', tag='table'):
    columns: list[Length]
    rows: Count
    cells: list[Cell]
    borders: Borders | None = None
    cell_margin: Length | None = None
    """The space on the sides of the content of the cells; Word puts a margin
    of its own without it."""


class PageBreak(msgspec.Struct, tag_field='type', tag='page_break'):
    pass


Block = Paragraph | Table | PageBreak


class WordExport(msgspec.Struct):
    title: str
    blocks: list[Block]
    margins: PageMargins = msgspec.field(default_factory=PageMargins)


# The fonts the files of the application provide, embedded in the Word files
# that use them, as the computer reading a file may lack them.
_EMBEDDED_FONTS = {PRINT_FONT_FAMILY: PRINT_FONT_FILES}

_PAGE_WIDTH = Mm(210)
_PAGE_HEIGHT = Mm(297)
_CELL_MARGIN = Mm(1.9)
_TAB_STOP_INSET = Mm(1)
# The width of a no-break space in Arial, in ems.
_NO_BREAK_SPACE_WIDTH = 0.278

_ALIGNMENTS = {
    'left': WD_ALIGN_PARAGRAPH.LEFT,
    'center': WD_ALIGN_PARAGRAPH.CENTER,
    'right': WD_ALIGN_PARAGRAPH.RIGHT,
    'justify': WD_ALIGN_PARAGRAPH.JUSTIFY,
}
_VERTICAL_ALIGNMENTS = {
    'top': WD_ALIGN_VERTICAL.TOP,
    'center': WD_ALIGN_VERTICAL.CENTER,
    'bottom': WD_ALIGN_VERTICAL.BOTTOM,
}
_SIDES = ('top', 'left', 'bottom', 'right')
_TABLE_CELL_PROPERTIES_AFTER_BORDERS = (
    'w:shd',
    'w:noWrap',
    'w:tcMar',
    'w:textDirection',
    'w:tcFitText',
    'w:vAlign',
    'w:hideMark',
)
_TABLE_PROPERTIES_AFTER_BORDERS = (
    'w:shd',
    'w:tblLayout',
    'w:tblCellMar',
    'w:tblLook',
    'w:tblCaption',
    'w:tblDescription',
)


def decode_word_export(content: bytes) -> WordExport:
    """Decode and validate the layout sent by the browser.

    :raises msgspec.ValidationError: when the content is not a valid layout.
    """
    return msgspec.json.decode(content, type=WordExport)


def build_word_document(export: WordExport) -> bytes:
    """The .docx file of an exported document."""
    document = new_document()
    section = document.sections[0]
    section.page_width = _PAGE_WIDTH
    section.page_height = _PAGE_HEIGHT
    margins = export.margins
    section.top_margin = Pt(margins.top)
    section.right_margin = Pt(margins.right)
    section.bottom_margin = Pt(margins.bottom)
    section.left_margin = Pt(margins.left)
    normal = document.styles['Normal']
    normal.font.name = 'Arial'
    normal.font.size = Pt(10)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.line_spacing = 1.0
    document.core_properties.title = export.title
    writer = _Writer(document)
    writer.write_blocks(
        document, export.blocks, Emu(_PAGE_WIDTH - Pt(margins.left) - Pt(margins.right))
    )
    for family in sorted(writer.fonts & _EMBEDDED_FONTS.keys()):
        embed_font(document, family, _EMBEDDED_FONTS[family])
    output = BytesIO()
    document.save(output)
    return output.getvalue()


class _Writer:
    def __init__(self, document: Document):
        self.document = document
        self.fonts: set[str] = set()
        """The font families of the runs written."""

    def write_blocks(
        self, container: Document | _Cell, blocks: list[Block], width: Emu
    ) -> None:
        """Write the blocks at the end of a container ``width`` wide.

        A page break moves the next block to a new page; the breaks with no
        block after them are dropped.
        """
        in_cell = isinstance(container, _Cell)
        break_pending = False
        previous: Block | None = None
        # The empty paragraph the next paragraph fills: a new cell holds one,
        # and Word ends every table in a cell with one.
        empty_paragraph: WordParagraph | None = (
            container.paragraphs[0] if isinstance(container, _Cell) else None
        )
        for block in blocks:
            if isinstance(block, PageBreak):
                break_pending = previous is not None
                continue
            if isinstance(block, Paragraph):
                if empty_paragraph is not None:
                    paragraph = empty_paragraph
                    empty_paragraph = None
                else:
                    paragraph = container.add_paragraph()
                self.write_paragraph(paragraph, block, width)
                if break_pending:
                    paragraph.paragraph_format.page_break_before = True
            else:
                if isinstance(previous, Table) and not in_cell:
                    # Two tables in a row would be joined into one.
                    _spacer(container.add_paragraph())
                table = self.write_table(container, block, width)
                if table is not None and break_pending:
                    table.cell(0, 0).paragraphs[
                        0
                    ].paragraph_format.page_break_before = True
                if isinstance(container, _Cell) and table is not None:
                    if previous is None and empty_paragraph is not None:
                        # The table opens the cell.
                        empty_paragraph._p.getparent().remove(empty_paragraph._p)
                    empty_paragraph = container.paragraphs[-1]
                    _spacer(empty_paragraph)
            break_pending = False
            previous = block

    def write_paragraph(
        self, paragraph: WordParagraph, block: Paragraph, width: Emu
    ) -> None:
        paragraph.paragraph_format.line_spacing = None
        paragraph.alignment = _ALIGNMENTS[block.align]
        if block.space_before:
            paragraph.paragraph_format.space_before = Pt(min(block.space_before, 24))
        runs = _trimmed(block.runs)
        for run in runs:
            if isinstance(run, TextRun):
                self.write_text_run(paragraph.add_run(run.text), run)
                self.use_font(paragraph.runs[-1], run.font)
            elif isinstance(run, BreakRun):
                paragraph.add_run().add_break()
            elif isinstance(run, BlankRun):
                self.write_blank_run(paragraph, run, width)
                self.use_font(paragraph.runs[-1], run.font)
            else:
                self.write_image_run(paragraph.add_run(), run, width)

    def use_font(self, word_run: WordRun, family: str | None) -> None:
        if family:
            word_run.font.name = family
            self.fonts.add(family)

    @staticmethod
    def write_text_run(word_run: WordRun, run: TextRun) -> None:
        word_run.bold = run.bold or None
        word_run.italic = run.italic or None
        if run.underline:
            word_run.underline = WD_UNDERLINE.SINGLE
        if run.strike:
            word_run.font.strike = True
        if run.size:
            word_run.font.size = Pt(run.size)
        if run.color and run.color.upper() != '000000':
            word_run.font.color.rgb = RGBColor.from_string(run.color.upper())

    @staticmethod
    def write_blank_run(paragraph: WordParagraph, run: BlankRun, width: Emu) -> None:
        if run.width is None:
            # A tab to a stop at the right edge, with a dotted leader, draws
            # the line to the end of the paragraph however long its text is.
            # LibreOffice draws no leader for a tab ending a paragraph in a
            # table cell, so a zero-width space follows it.
            paragraph.paragraph_format.tab_stops.add_tab_stop(
                Emu(max(0, width - _TAB_STOP_INSET)),
                WD_TAB_ALIGNMENT.LEFT,
                WD_TAB_LEADER.DOTS,
            )
            word_run = paragraph.add_run('\t\u200b')
        else:
            size = run.size or 10
            count = max(4, min(200, round(run.width / (size * _NO_BREAK_SPACE_WIDTH))))
            word_run = paragraph.add_run('\u00a0' * count)
            word_run.underline = WD_UNDERLINE.DOTTED
        if run.size:
            word_run.font.size = Pt(run.size)

    @staticmethod
    def write_image_run(word_run: WordRun, run: ImageRun, width: Emu) -> None:
        image = _decode_image(run.src)
        if image is None or not run.width or not run.height:
            return
        scale = min(1.0, width / Pt(run.width))
        word_run.add_picture(
            image,
            width=Emu(int(Pt(run.width) * scale)),
            height=Emu(int(Pt(run.height) * scale)),
        )

    def write_table(
        self, container: Document | _Cell, block: Table, width: Emu
    ) -> WordTable | None:
        total = sum(block.columns)
        if not block.columns or not total:
            return None
        column_widths = [Emu(int(width * column / total)) for column in block.columns]
        table = container.add_table(block.rows, len(block.columns))
        table.autofit = False
        for column, column_width in zip(table.columns, column_widths, strict=True):
            column.width = column_width
        if block.borders:
            _set_borders(
                table._tbl.tblPr,
                'w:tblBorders',
                block.borders,
                _TABLE_PROPERTIES_AFTER_BORDERS,
            )
        cell_margin: WordLength = _CELL_MARGIN
        if block.cell_margin is not None:
            cell_margin = Emu(Pt(min(block.cell_margin, 24)))
            _set_side_cell_margins(table, cell_margin)
        if cell_margin:
            _indent(table, cell_margin)
        placed: list[tuple[_Cell, Cell, Emu]] = []
        for cell in block.cells:
            last_row = cell.row + cell.row_span - 1
            last_column = cell.column + cell.column_span - 1
            if last_row >= block.rows or last_column >= len(block.columns):
                continue
            word_cell = table.cell(cell.row, cell.column)
            if last_row > cell.row or last_column > cell.column:
                try:
                    word_cell = word_cell.merge(table.cell(last_row, last_column))
                except InvalidSpanError:
                    continue
            cell_width = Emu(sum(column_widths[cell.column : last_column + 1]))
            placed.append((word_cell, cell, cell_width))
        for row in table.rows:
            for word_cell, column_width in zip(row.cells, column_widths, strict=False):
                word_cell.width = column_width
        for word_cell, cell, cell_width in placed:
            word_cell.width = cell_width
            word_cell.vertical_alignment = _VERTICAL_ALIGNMENTS[cell.vertical_align]
            if cell.borders:
                _set_borders(
                    word_cell._tc.get_or_add_tcPr(),
                    'w:tcBorders',
                    cell.borders,
                    _TABLE_CELL_PROPERTIES_AFTER_BORDERS,
                )
            if cell.fill:
                _set_fill(word_cell, cell.fill)
            if cell.padding:
                _set_vertical_padding(word_cell, min(cell.padding, 24))
            self.write_blocks(
                word_cell, cell.blocks, Emu(max(0, cell_width - 2 * cell_margin))
            )
            if cell.space_before and word_cell.paragraphs:
                word_cell.paragraphs[0].paragraph_format.space_before = Pt(
                    min(cell.space_before, 24)
                )
        return table


def _spacer(paragraph: WordParagraph) -> None:
    """Make an empty paragraph a thin space between two tables."""
    paragraph.paragraph_format.line_spacing = Pt(4)


def _trimmed(runs: list[Run]) -> list[Run]:
    """The runs without the spaces that open and close the paragraph."""
    runs = list(runs)
    if runs and isinstance(first := runs[0], TextRun):
        runs[0] = msgspec.structs.replace(first, text=first.text.lstrip())
    if runs and isinstance(last := runs[-1], TextRun):
        runs[-1] = msgspec.structs.replace(last, text=last.text.rstrip())
    return [run for run in runs if not (isinstance(run, TextRun) and not run.text)]


def _decode_image(src: str) -> BytesIO | None:
    """The image of a ``data:`` URL, when it is one Word can show."""
    header, _, data = src.partition(',')
    if not header.startswith('data:image/') or not header.endswith(';base64'):
        return None
    try:
        content = base64.b64decode(data, validate=True)
        with Image.open(BytesIO(content)) as image:
            if image.format not in ('PNG', 'JPEG', 'GIF', 'BMP'):
                return None
    except (binascii.Error, UnidentifiedImageError, OSError):
        return None
    return BytesIO(content)


def _set_borders(
    properties: BaseOxmlElement,
    tag: str,
    borders: Borders,
    successors: tuple[str, ...],
) -> None:
    element = OxmlElement(tag)
    for side in _SIDES:
        border: Border | None = getattr(borders, side)
        if border is None or not border.width:
            continue
        side_element = OxmlElement(f'w:{side}')
        side_element.set(qn('w:val'), border.style)
        # Word sizes the borders in eighths of a point, from 2 to 96.
        side_element.set(qn('w:sz'), str(max(2, min(96, round(border.width * 8)))))
        side_element.set(qn('w:space'), '0')
        side_element.set(qn('w:color'), border.color.upper())
        element.append(side_element)
    properties.insert_element_before(element, *successors)


def _set_fill(cell: _Cell, color: str) -> None:
    shading = OxmlElement('w:shd')
    shading.set(qn('w:val'), 'clear')
    shading.set(qn('w:color'), 'auto')
    shading.set(qn('w:fill'), color.upper())
    cell._tc.get_or_add_tcPr().insert_element_before(
        shading, *_TABLE_CELL_PROPERTIES_AFTER_BORDERS[1:]
    )


def _set_vertical_padding(cell: _Cell, padding: float) -> None:
    margins = OxmlElement('w:tcMar')
    for side in ('top', 'bottom'):
        margin = OxmlElement(f'w:{side}')
        # Word sizes the cell margins in twentieths of a point.
        margin.set(qn('w:w'), str(round(padding * 20)))
        margin.set(qn('w:type'), 'dxa')
        margins.append(margin)
    cell._tc.get_or_add_tcPr().insert_element_before(
        margins, *_TABLE_CELL_PROPERTIES_AFTER_BORDERS[3:]
    )


def _indent(table: WordTable, indent: WordLength) -> None:
    """Indent a table so that its border, rather than the text of its cells,
    lines up with the text around it: Word draws a table its cell margin to
    the left of its indent."""
    element = OxmlElement('w:tblInd')
    element.set(qn('w:w'), str(round(indent.pt * 20)))
    element.set(qn('w:type'), 'dxa')
    table._tbl.tblPr.insert_element_before(
        element,
        'w:tblBorders',
        'w:shd',
        'w:tblLayout',
        'w:tblCellMar',
        'w:tblLook',
        'w:tblCaption',
        'w:tblDescription',
    )


def _set_side_cell_margins(table: WordTable, margin_width: WordLength) -> None:
    margins = OxmlElement('w:tblCellMar')
    for side in ('left', 'right'):
        margin = OxmlElement(f'w:{side}')
        margin.set(qn('w:w'), str(round(margin_width.pt * 20)))
        margin.set(qn('w:type'), 'dxa')
        margins.append(margin)
    table._tbl.tblPr.insert_element_before(
        margins, 'w:tblLook', 'w:tblCaption', 'w:tblDescription'
    )
