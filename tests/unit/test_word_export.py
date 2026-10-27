import base64
import json
from io import BytesIO
from typing import Any

import msgspec
import pytest
from docx import Document as open_document
from docx.document import Document
from docx.enum.text import WD_TAB_LEADER
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.oxml import parse_xml
from docx.oxml.ns import qn
from docx.shared import Pt
from PIL import Image

from data.print_documents.fonts import PRINT_FONT_FAMILY, PRINT_FONT_FILES
from data.print_documents.word_font_embedding import obfuscated
from data.print_documents.word_export import (
    build_word_document,
    decode_word_export,
)


def _export(*blocks: dict[str, Any]) -> Document:
    content = json.dumps({'title': 'Test', 'blocks': list(blocks)}).encode()
    return open_document(BytesIO(build_word_document(decode_word_export(content))))


def _paragraph(*runs: dict[str, Any], **properties: Any) -> dict[str, Any]:
    return {'type': 'paragraph', 'runs': list(runs)} | properties


def _text(text: str, **properties: Any) -> dict[str, Any]:
    return {'type': 'text', 'text': text} | properties


def _cell(row: int, column: int, text: str, **properties: Any) -> dict[str, Any]:
    return {
        'row': row,
        'column': column,
        'blocks': [_paragraph(_text(text))],
    } | properties


def _png_data_url() -> str:
    output = BytesIO()
    Image.new('RGB', (20, 10), 'red').save(output, format='PNG')
    return 'data:image/png;base64,' + base64.b64encode(output.getvalue()).decode()


@pytest.mark.unit
class TestWordExport:
    def test_title(self) -> None:
        document = _export(_paragraph(_text('Body')))
        assert document.core_properties.title == 'Test'

    def test_text_runs(self) -> None:
        document = _export(
            _paragraph(
                _text('  Name: '),
                _text('Doe', bold=True, size=12, color='FF0000'),
                _text(' struck ', strike=True),
                align='center',
            )
        )
        paragraph = document.paragraphs[0]
        assert paragraph.text == 'Name: Doe struck'
        assert paragraph.alignment == 1
        name, value, struck = paragraph.runs
        assert not name.bold
        assert value.bold
        assert value.font.size == Pt(12)
        assert str(value.font.color.rgb) == 'FF0000'
        assert struck.font.strike

    def test_page_margins(self) -> None:
        content = json.dumps(
            {
                'title': 'Test',
                'blocks': [_paragraph(_text('Body'))],
                'margins': {
                    'top': 14.17,
                    'right': 14.17,
                    'bottom': 14.17,
                    'left': 28.35,
                },
            }
        ).encode()
        document = open_document(
            BytesIO(build_word_document(decode_word_export(content)))
        )
        section = document.sections[0]
        assert section.left_margin is not None
        assert section.right_margin is not None
        assert round(section.left_margin.mm) == 10
        assert round(section.right_margin.mm) == 5

    def test_default_page_margins(self) -> None:
        section = _export(_paragraph(_text('Body'))).sections[0]
        assert section.top_margin is not None
        assert section.left_margin is not None
        assert round(section.top_margin.mm) == round(section.left_margin.mm) == 10

    def test_line_break(self) -> None:
        document = _export(
            _paragraph(_text('Formulaire'), {'type': 'break'}, _text('Rapport'))
        )
        assert document.paragraphs[0].text == 'Formulaire\nRapport'

    def test_blank_to_the_end_of_the_paragraph(self) -> None:
        document = _export(_paragraph(_text('Address: '), {'type': 'blank'}))
        paragraph = document.paragraphs[0]
        assert paragraph.runs[-1].text == '\t\u200b'
        (tab_stop,) = paragraph.paragraph_format.tab_stops
        assert tab_stop.leader == WD_TAB_LEADER.DOTS

    def test_blank_of_a_width(self) -> None:
        document = _export(
            _paragraph(_text('Adults = '), {'type': 'blank', 'width': 28, 'size': 10})
        )
        blank = document.paragraphs[0].runs[-1]
        assert set(blank.text) == {'\u00a0'}
        assert len(blank.text) == 10
        assert blank.font.underline is not None

    def test_page_breaks(self) -> None:
        document = _export(
            {'type': 'page_break'},
            _paragraph(_text('First')),
            {'type': 'page_break'},
            _paragraph(_text('Second')),
            {'type': 'page_break'},
        )
        first, second = document.paragraphs
        assert not first.paragraph_format.page_break_before
        assert second.paragraph_format.page_break_before

    def test_table_with_merged_cells(self) -> None:
        document = _export(
            {
                'type': 'table',
                'columns': [1, 2, 1],
                'rows': 2,
                'cells': [
                    _cell(0, 0, 'Letter', row_span=2),
                    _cell(0, 1, 'Club', column_span=2),
                    _cell(1, 1, 'Town'),
                    _cell(1, 2, 'Code'),
                ],
            }
        )
        table = document.tables[0]
        assert [cell.text for cell in table.rows[0].cells] == [
            'Letter',
            'Club',
            'Club',
        ]
        assert [cell.text for cell in table.rows[1].cells] == [
            'Letter',
            'Town',
            'Code',
        ]
        widths = [column.width for column in table.columns]
        assert widths[1] == 2 * widths[0]
        assert widths[0] == widths[2]

    def test_cells_out_of_the_table_are_dropped(self) -> None:
        document = _export(
            {
                'type': 'table',
                'columns': [1],
                'rows': 1,
                'cells': [_cell(0, 0, 'In'), _cell(1, 0, 'Out'), _cell(0, 1, 'Out')],
            }
        )
        assert [cell.text for cell in document.tables[0].rows[0].cells] == ['In']

    def test_cell_borders_and_fill(self) -> None:
        document = _export(
            {
                'type': 'table',
                'columns': [1],
                'rows': 1,
                'cells': [
                    _cell(
                        0,
                        0,
                        'Boxed',
                        borders={'top': {'width': 1}, 'left': {'width': 0}},
                        fill='E0E0E0',
                        padding=6,
                    )
                ],
            }
        )
        properties = document.tables[0].cell(0, 0)._tc.tcPr
        assert properties is not None
        borders = properties.find(qn('w:tcBorders'))
        assert [side.tag for side in borders] == [qn('w:top')]
        assert borders[0].get(qn('w:sz')) == '8'
        assert properties.find(qn('w:shd')).get(qn('w:fill')) == 'E0E0E0'
        margins = properties.find(qn('w:tcMar'))
        assert [margin.get(qn('w:w')) for margin in margins] == ['120', '120']

    def test_table_without_cell_margins(self) -> None:
        document = _export(
            {
                'type': 'table',
                'columns': [1, 1],
                'rows': 1,
                'cells': [_cell(0, 0, 'Name'), _cell(0, 1, 'Doe')],
                'cell_margin': 0,
            }
        )
        margins = document.tables[0]._tbl.tblPr.find(qn('w:tblCellMar'))
        assert [(margin.tag, margin.get(qn('w:w'))) for margin in margins] == [
            (qn('w:left'), '0'),
            (qn('w:right'), '0'),
        ]
        assert document.tables[0]._tbl.tblPr.find(qn('w:tblInd')) is None

    def test_cell_margins(self) -> None:
        document = _export(
            {
                'type': 'table',
                'columns': [1],
                'rows': 1,
                'cells': [_cell(0, 0, 'Cell')],
                'cell_margin': 3,
            }
        )
        properties = document.tables[0]._tbl.tblPr
        margins = properties.find(qn('w:tblCellMar'))
        assert [margin.get(qn('w:w')) for margin in margins] == ['60', '60']
        indent = properties.find(qn('w:tblInd'))
        assert indent is not None
        assert indent.get(qn('w:w')) == '60'

    def test_bordered_table_lines_up_with_the_text(self) -> None:
        document = _export(
            {'type': 'table', 'columns': [1], 'rows': 1, 'cells': [_cell(0, 0, 'Box')]}
        )
        indent = document.tables[0]._tbl.tblPr.find(qn('w:tblInd'))
        assert indent is not None
        assert indent.get(qn('w:w')) == '108'

    def test_tables_in_a_row_stay_apart(self) -> None:
        table = {
            'type': 'table',
            'columns': [1],
            'rows': 1,
            'cells': [_cell(0, 0, 'Cell')],
        }
        document = _export(table, table)
        body = [element.tag for element in document.element.body]
        assert body[:3] == [qn('w:tbl'), qn('w:p'), qn('w:tbl')]

    def test_nested_table(self) -> None:
        document = _export(
            {
                'type': 'table',
                'columns': [1],
                'rows': 1,
                'cells': [
                    {
                        'row': 0,
                        'column': 0,
                        'blocks': [
                            _paragraph(_text('Section')),
                            {
                                'type': 'table',
                                'columns': [1, 1],
                                'rows': 1,
                                'cells': [_cell(0, 0, 'Name'), _cell(0, 1, 'Doe')],
                            },
                        ],
                    }
                ],
            }
        )
        cell = document.tables[0].cell(0, 0)
        assert cell.paragraphs[0].text == 'Section'
        assert [element.tag for element in cell._tc][1:] == [
            qn('w:p'),
            qn('w:tbl'),
            qn('w:p'),
        ]
        assert [inner.text for inner in cell.tables[0].rows[0].cells] == [
            'Name',
            'Doe',
        ]

    def test_table_opening_a_cell(self) -> None:
        inner = {
            'type': 'table',
            'columns': [1],
            'rows': 1,
            'cells': [_cell(0, 0, 'Inner')],
        }
        document = _export(
            {
                'type': 'table',
                'columns': [1],
                'rows': 1,
                'cells': [
                    {
                        'row': 0,
                        'column': 0,
                        'blocks': [inner, inner, _paragraph(_text('After'))],
                    }
                ],
            }
        )
        cell = document.tables[0].cell(0, 0)
        assert [element.tag for element in cell._tc][1:] == [
            qn('w:tbl'),
            qn('w:p'),
            qn('w:tbl'),
            qn('w:p'),
        ]
        assert cell.paragraphs[-1].text == 'After'

    def test_image(self) -> None:
        document = _export(
            _paragraph(
                {'type': 'image', 'src': _png_data_url(), 'width': 40, 'height': 20}
            )
        )
        assert len(document.inline_shapes) == 1
        assert document.inline_shapes[0].width == Pt(40)

    def test_image_wider_than_the_page_is_scaled_down(self) -> None:
        document = _export(
            _paragraph(
                {'type': 'image', 'src': _png_data_url(), 'width': 2000, 'height': 1000}
            )
        )
        shape = document.inline_shapes[0]
        assert shape.width < Pt(2000)
        assert abs(shape.width - 2 * shape.height) <= 2

    @pytest.mark.parametrize(
        'src',
        [
            'https://example.com/logo.png',
            'data:image/png;base64,not-base64',
            'data:image/svg+xml;base64,'
            + base64.b64encode(b'<svg xmlns="http://www.w3.org/2000/svg"/>').decode(),
        ],
    )
    def test_images_word_cannot_show_are_dropped(self, src: str) -> None:
        document = _export(
            _paragraph(
                _text('Logo'), {'type': 'image', 'src': src, 'width': 4, 'height': 4}
            )
        )
        assert len(document.inline_shapes) == 0
        assert document.paragraphs[0].text == 'Logo'

    def test_print_font_is_embedded(self) -> None:
        document = _export(_paragraph(_text('Formulaire', font=PRINT_FONT_FAMILY)))
        assert document.paragraphs[0].runs[0].font.name == PRINT_FONT_FAMILY
        font_table = document.part.part_related_by(RELATIONSHIP_TYPE.FONT_TABLE)
        (font,) = [
            font
            for font in parse_xml(font_table.blob).findall(qn('w:font'))
            if font.get(qn('w:name')) == PRINT_FONT_FAMILY
        ]
        embedded = list(font)
        assert [element.tag for element in embedded] == [
            qn('w:embedRegular'),
            qn('w:embedBold'),
            qn('w:embedItalic'),
            qn('w:embedBoldItalic'),
        ]
        regular = font_table.related_parts[embedded[0].get(qn('r:id'))]
        assert regular.content_type == (
            'application/vnd.openxmlformats-officedocument.obfuscatedFont'
        )
        key = embedded[0].get(qn('w:fontKey'))
        assert obfuscated(regular.blob, key) == (
            PRINT_FONT_FILES[False, False].read_bytes()
        )
        assert document.settings.element.find(qn('w:embedTrueTypeFonts')) is not None

    def test_other_fonts_are_named_only(self) -> None:
        document = _export(_paragraph(_text('Name', font='Arial')))
        assert document.paragraphs[0].runs[0].font.name == 'Arial'
        assert document.settings.element.find(qn('w:embedTrueTypeFonts')) is None

    def test_obfuscation(self) -> None:
        font = bytes(range(64))
        key = '{0F1E2D3C-4B5A-6978-8796-A5B4C3D2E1F0}'
        assert obfuscated(font, key)[32:] == font[32:]
        masks = [
            bytes(
                a ^ b
                for a, b in zip(
                    obfuscated(font, key)[start : start + 16],
                    font[start : start + 16],
                    strict=True,
                )
            )
            for start in (0, 16)
        ]
        assert masks[0] == masks[1] == bytes.fromhex('F0E1D2C3B4A5968778695A4B3C2D1E0F')
        assert obfuscated(obfuscated(font, key), key) == font

    @pytest.mark.parametrize(
        'block',
        [
            {'type': 'unknown'},
            _paragraph(_text('Red', color='red')),
            _paragraph(_text('Huge', size=1000)),
            {'type': 'table', 'columns': [1], 'rows': 0, 'cells': []},
        ],
    )
    def test_invalid_layouts_are_refused(self, block: dict[str, Any]) -> None:
        content = json.dumps({'title': 'Test', 'blocks': [block]}).encode()
        with pytest.raises(msgspec.ValidationError):
            decode_word_export(content)
