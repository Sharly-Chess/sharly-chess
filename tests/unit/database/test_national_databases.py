"""The rating lists of the federations converted locally: reading each
federation's file, and searching the shared player table."""

import sqlite3
from datetime import date
from pathlib import Path

import pytest

from database.sqlite.national.cfc_database import CfcDatabase
from database.sqlite.national.cfr_database import CfrDatabase
from database.sqlite.national.fsi_database import FsiDatabase
from database.sqlite.national.knsb_database import KnsbDatabase
from database.sqlite.national.national_database import (
    NationalPlayerDatabase,
    NationalPlayerRow,
)
from database.sqlite.national.ssl_database import SslDatabase
from database.sqlite.sqlite_database import SQLiteDatabase


def _write(path: Path, content: str, encoding: str = 'utf-8') -> Path:
    path.write_text(content, encoding=encoding, newline='')
    return path


@pytest.mark.unit
def test_knsb_reads_the_three_rate_of_play_lists(tmp_path: Path) -> None:
    header = 'Relatienummer;Naam;Titel;FED;Rating;Nv;Geboren;S\r\n'
    _write(
        tmp_path / 'RAPID.csv',
        header + '9055882;Van der Aa, Quinten;;NED;1658;42;2012;\r\n',
        'cp1252',
    )
    _write(
        tmp_path / 'SNEL.csv',
        header + '9055882;Van der Aa, Quinten;;NED;1750;44;2012;\r\n',
        'cp1252',
    )
    file = _write(
        tmp_path / 'KLASSIEK.csv',
        header
        + '9055882;Van der Aa, Quinten;;NED;1679;24;2012;\r\n'
        + '8538442;Aarten, Marjolein;WFM;NED;;100;1981;w\r\n'
        + '8888715;Van der Aart, André;;GER;1544;24;1958;\r\n',
        'cp1252',
    )
    rows = list(KnsbDatabase()._read_players(file))
    assert [row.national_id for row in rows] == ['9055882', '8538442', '8888715']
    quinten, marjolein, andre = rows
    assert (quinten.last_name, quinten.first_name) == ('Van der Aa', 'Quinten')
    assert (quinten.standard_rating, quinten.rapid_rating, quinten.blitz_rating) == (
        1679,
        1658,
        1750,
    )
    assert (quinten.gender, quinten.year_of_birth) == ('M', 2012)
    assert (marjolein.gender, marjolein.title, marjolein.standard_rating) == (
        'F',
        'WFM',
        None,
    )
    assert (marjolein.rapid_rating, marjolein.blitz_rating) == (None, None)
    assert (andre.first_name, andre.federation) == ('André', 'GER')


@pytest.mark.unit
def test_fsi_reads_the_allin_export(tmp_path: Path) -> None:
    header = (
        'Nominativo;Elo;Categoria;K Italia;Federazione;Anno Nascita;ID FSI;Sex;FIN;'
        'Elo FIDE Std;K Elo FIDE Std;Games Elo FIDE Std;Elo FIDE Rapid;K Elo FIDE Rapid;'
        'Games Elo FIDE Rapid;Elo FIDE Blitz;K Elo FIDE Blitz;Games Elo FIDE Blitz;'
        'Attivo;Data Nascita;Provincia Residenza;Regione Residenza;Provincia Tesseramento;'
        'Regione Tesseramento;Cittadinanza;Data Ultimo Torneo;Circolo Attuale;'
        'Circolo Anno Precedente;Circolo Due Anni Precedenti;Anno Ultimo Tesseramento;'
        'Tipo Tessera;Numero Tessera;\r\n'
    )
    file = _write(
        tmp_path / 'allin.csv',
        header
        + 'AAB Manfred;0;;0;;0;0;m;;0;0;0;0;0;0;0;0;0;;;;;;;;;0;0;0;0;;0;\r\n'
        + 'AARON ALEXANDER;1399;NC;30;ITA;2010;169546;M;2895027;1512;40;12;0;0;0;1600;20;3;'
        'Y;15-02-2010;MI;LOM;;;ITA;06-03-2022;0;0;0;2023;;0;\r\n'
        + 'BAN NICCOLÒ;0;;0;;2004;195382;M;;0;0;0;0;0;0;0;0;0;;;;;;;;;0;0;0;0;;0;\r\n'
        + 'DE ROSSI Anna Maria;1800;1N;20;ITA;1990;100;F;;0;0;0;0;0;0;0;0;0;;;;;;;;;0;0;0;0;;0;\r\n',
    )
    rows = list(FsiDatabase()._read_players(file))
    assert [row.national_id for row in rows] == ['169546', '195382', '100']
    alexander, niccolo, anna = rows
    assert (alexander.last_name, alexander.first_name) == ('AARON', 'ALEXANDER')
    assert alexander.date_of_birth is not None
    assert (alexander.date_of_birth.isoformat(), alexander.year_of_birth) == (
        '2010-02-15',
        2010,
    )
    assert (alexander.standard_rating, alexander.fide_id) == (1399, 2895027)
    assert (
        alexander.fide_standard_rating,
        alexander.fide_rapid_rating,
        alexander.fide_blitz_rating,
    ) == (1512, None, 1600)
    assert (niccolo.last_name, niccolo.first_name, niccolo.year_of_birth) == (
        'BAN',
        'NICCOLÒ',
        2004,
    )
    assert (anna.last_name, anna.first_name, anna.gender) == (
        'DE ROSSI',
        'Anna Maria',
        'F',
    )


@pytest.mark.unit
def test_cfr_reads_the_swiss_manager_exports(tmp_path: Path) -> None:
    header = (
        'ID_No,Name,Sex,Fed,Clubnumber,ClubName,Birthday,Rtg_Nat,Fide_No,Rtg_Int\r\n'
    )
    _write(
        tmp_path / 'smanager_rapid.csv',
        header
        + '10,Несытова Анастасия Максимовна,f,RUS,23,Краснодарский край,1997,2250,4164970,2201\r\n',
    )
    _write(tmp_path / 'smanager_blitz.csv', header)
    file = _write(
        tmp_path / 'smanager_standard.csv',
        '﻿'
        + header
        + '1,Переверткин Владимир Викторович,,RUS,77,Москва,1967,,4112342,2372\r\n'
        + '10,Несытова Анастасия Максимовна,f,RUS,23,Краснодарский край,1997,2213,4164970,2201\r\n',
    )
    rows = list(CfrDatabase()._read_players(file))
    vladimir, anastasia = rows
    assert (vladimir.national_id, vladimir.last_name, vladimir.first_name) == (
        '1',
        'Переверткин',
        'Владимир Викторович',
    )
    assert (vladimir.gender, vladimir.standard_rating, vladimir.fide_id) == (
        'M',
        None,
        4112342,
    )
    assert vladimir.fide_standard_rating == 2372
    assert (anastasia.gender, anastasia.club) == ('F', 'Краснодарский край')
    assert (
        anastasia.standard_rating,
        anastasia.rapid_rating,
        anastasia.blitz_rating,
    ) == (2213, 2250, None)


@pytest.mark.unit
def test_ssl_reads_the_selo_export(tmp_path: Path) -> None:
    file = _write(
        tmp_path / 'selo.csv',
        'Id;Sukunimi;Etunimi;Seura;Selo;Selopelejä;Ikäryhmä;Maa;Lisenssi;\r\n'
        '9415;Shirov;Aleksei;TammerSh;2739;30;S50;ESP;;\r\n'
        '5239;Nybäck;Tomi;Aatos;2601;2183;;FIN;L;\r\n',
        'cp1252',
    )
    rows = list(SslDatabase()._read_players(file))
    shirov, nyback = rows
    assert (shirov.national_id, shirov.federation, shirov.club) == (
        '9415',
        'ESP',
        'TammerSh',
    )
    assert (nyback.last_name, nyback.standard_rating, nyback.gender) == (
        'Nybäck',
        2601,
        '',
    )


@pytest.mark.unit
def test_cfc_reads_the_td_list_despite_its_quoting(tmp_path: Path) -> None:
    file = _write(
        tmp_path / 'tdlist.txt',
        '"CFC#","Expiry","Last","First","Prov","City","Rating","High","Active Rtg",'
        '"Active High","FIDE Number","FIDE Rating"\r\n'
        '191362,2025-12-18,"---","---","ON",Scarborough",0,0,1092,1,,0\r\n'
        '171117,2019-08-20," Lancman","Kyle","US",New York",1991,1991,1898,1806,75003732,2100\r\n'
        '165306,2017-10-15," Ngo","Ethan","AB",Calgary",0,0,208,4,,0\r\n',
    )
    rows = list(CfcDatabase()._read_players(file))
    assert [row.national_id for row in rows] == ['171117', '165306']
    kyle, ethan = rows
    assert (kyle.last_name, kyle.first_name, kyle.club) == (
        'Lancman',
        'Kyle',
        'New York',
    )
    assert (kyle.standard_rating, kyle.rapid_rating, kyle.blitz_rating) == (
        1991,
        1898,
        1898,
    )
    assert (kyle.fide_id, kyle.fide_standard_rating) == (75003732, 2100)
    assert (ethan.standard_rating, ethan.rapid_rating) == (None, 208)


@pytest.mark.unit
def test_last_name_split_of_capitalised_names() -> None:
    split = NationalPlayerDatabase._split_upper_last_name
    assert split('DE ROSSI Anna Maria') == ('DE ROSSI', 'Anna Maria')
    assert split('AARON ALEXANDER') == ('AARON', 'ALEXANDER')
    assert split('Rossi Mario') == ('Rossi', 'Mario')
    assert split('ROSSI') == ('ROSSI', '')
    assert split('') == ('', '')


@pytest.fixture
def national_database() -> KnsbDatabase:
    """A KNSB database over an in-memory copy of the shared player table."""
    connection = sqlite3.connect(':memory:')
    connection.row_factory = sqlite3.Row
    database = KnsbDatabase()
    connection.executescript(database._schema)
    rows = [
        NationalPlayerRow(
            '1',
            'Dupont',
            'Alice',
            2000,
            gender='F',
            federation='NED',
            club='Amsterdam',
            standard_rating=1800,
            fide_id=1001,
        ),
        NationalPlayerRow(
            '2',
            'Dupont',
            'Benoit',
            1992,
            gender='M',
            federation='NED',
            club='Utrecht',
            rapid_rating=1700,
        ),
        NationalPlayerRow(
            '3',
            'Dupont',
            'Christine',
            1960,
            gender='F',
            federation='BEL',
            club='Amsterdam',
        ),
        NationalPlayerRow('4', 'Martin', 'Dupont', 2012, gender='M', federation='NED'),
        NationalPlayerRow('12', 'Nybäck', 'Éva', 1986, gender='F', federation='NED'),
        NationalPlayerRow(
            '13', 'Бондаренко', 'Артём', 1990, gender='M', federation='RUS'
        ),
    ]
    connection.executemany(
        'INSERT INTO player VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)',
        [row.params for row in rows],
    )
    connection.commit()
    database.database = connection
    database.cursor = connection.cursor()
    return database


@pytest.mark.unit
def test_search_by_name_ranks_last_names_first(national_database: KnsbDatabase) -> None:
    result = national_database.search_player('dupont', 'NED', 0, None, {})
    assert [player.first_name for player in result] == [
        'Alice',
        'Benoit',
        'Christine',
        'Dupont',
    ]
    assert result[0].last_name == 'DUPONT'
    assert result[0].national_id == '1'
    assert result[0].national_source == 'knsb'
    assert result[0].ratings[1] == {'national': 1800}
    assert result[1].ratings[2] == {'national': 1700}


@pytest.mark.unit
def test_search_by_identifier(national_database: KnsbDatabase) -> None:
    assert [
        player.first_name
        for player in national_database.search_player('1', 'NED', 0, None, {})
    ] == ['Alice']
    assert [
        player.first_name
        for player in national_database.search_player('1001', 'NED', 0, None, {})
    ] == ['Alice']
    eva = national_database.get_stored_player_by_national_id('12')
    assert eva is not None and eva.first_name == 'Éva'
    assert national_database.get_stored_player_by_national_id('99') is None
    assert sorted(
        player.national_id or ''
        for player in national_database.get_stored_players_by_national_id(
            ['2', '3', '99']
        )
    ) == ['2', '3']


@pytest.mark.unit
def test_search_filters(national_database: KnsbDatabase) -> None:
    def first_names(filters: dict) -> list[str | None]:
        return [
            player.first_name
            for player in national_database.search_player(
                'dupont', 'NED', 0, None, filters
            )
        ]

    assert first_names({'federation_filter': 'BEL'}) == ['Christine']
    assert first_names({'gender_filter': 'F'}) == ['Alice', 'Christine']
    assert first_names({'club_filter': 'amster'}) == ['Alice', 'Christine']
    assert first_names({'year_of_birth_filter': [(1990, None)]}) == [
        'Alice',
        'Benoit',
        'Dupont',
    ]
    assert first_names({'year_of_birth_filter': [(None, 1970), (2010, 2020)]}) == [
        'Christine',
        'Dupont',
    ]


@pytest.mark.unit
def test_search_ignores_accents_and_case(national_database: KnsbDatabase) -> None:
    for string in ('nyback', 'NYBÄCK', 'Nybäck éva', 'eva'):
        assert [
            player.last_name
            for player in national_database.search_player(string, 'NED', 0, None, {})
        ] == ['NYBÄCK'], string


@pytest.mark.unit
def test_search_transliterates_cyrillic(national_database: KnsbDatabase) -> None:
    for string in ('bondarenko', 'Бондаренко', 'BONDARENKO artiom', 'артём'):
        assert [
            player.last_name
            for player in national_database.search_player(string, 'RUS', 0, None, {})
        ] == ['БОНДАРЕНКО'], string


@pytest.mark.unit
def test_search_pages(national_database: KnsbDatabase) -> None:
    assert [
        player.first_name
        for player in national_database.search_player('dupont', 'NED', 1, 2, {})
    ] == ['Christine', 'Dupont']


@pytest.mark.unit
def test_populate_stores_the_rows(tmp_path: Path) -> None:
    file = _write(
        tmp_path / 'selo.csv',
        'Id;Sukunimi;Etunimi;Seura;Selo;Selopelejä;Ikäryhmä;Maa;Lisenssi;\r\n'
        '9415;Shirov;Aleksei;TammerSh;2739;30;S50;ESP;;\r\n',
        'cp1252',
    )
    database_file = tmp_path / 'ssl.db'
    ssl_database = SslDatabase()
    SQLiteDatabase(database_file, write=True)._create(ssl_database._schema)
    with SQLiteDatabase(database_file, write=True) as database:
        assert ssl_database._populate_from_source_file(file, database)
        SslDatabase._create_indexes(database)
    with SQLiteDatabase(database_file) as database:
        database.execute('SELECT last_name, standard_rating FROM player')
        assert list(database.fetchall()) == [
            {'last_name': 'Shirov', 'standard_rating': 2739}
        ]


@pytest.mark.unit
def test_dsb_reads_the_players_with_their_club_and_keeps_the_active_membership(
    tmp_path: Path,
) -> None:
    from database.sqlite.national.dsb_database import DsbDatabase

    _write(
        tmp_path / 'vereine.csv',
        'ZPS-Nummer,Landesverband,UebergeordneterVerband,Vereinsname\r\n'
        '10614,C,C00,"FC ST.Pauli 1910 eV SAbt"\r\n'
        '20001,1,100,"Schachfreunde Süd"\r\n',
        'cp1252',
    )
    header = (
        'ID,ZPS,Mitgliedsnummer,Status,"Name,Vorname",Geschlecht,Spielberechtigung,'
        'Geburtsjahr,"Letzte Auswertung",DWZ,Index,FIDE-Elozahl,FIDE-Titel,FIDE-ID,'
        'FIDE-Land,Vorname,Nachname,FIDE-Frauentitel,FIDE-Elozahl-Schnellschach,'
        'FIDE-Elozahl-Blitz\r\n'
    )
    file = _write(
        tmp_path / 'spieler.csv',
        header
        + 'NU4073762,10614,1081,P,"Carlsen,Magnus",M,,1990,202517,2841,106,2823,GM,'
        '1503014,NOR,Magnus,Carlsen,,2803,2860\r\n'
        + 'NU0000001,20001,12,P,"Müller,Anna",W,,2005,,1650,20,,,,,Anna,Müller,WFM,,\r\n'
        + 'NU0000001,10614,13,A,"Müller,Anna",W,,2005,,1650,20,,,,,Anna,Müller,WFM,,\r\n'
        + 'NU0000002,99999,1,,"Doe,John",,,,,,,,,,,John,Doe,,,\r\n',
        'cp1252',
    )
    rows = {row.national_id: row for row in DsbDatabase()._read_players(file)}
    assert sorted(rows) == ['NU0000001', 'NU0000002', 'NU4073762']
    carlsen = rows['NU4073762']
    assert (carlsen.club, carlsen.federation, carlsen.title, carlsen.gender) == (
        'FC ST.Pauli 1910 eV SAbt',
        'NOR',
        'GM',
        'M',
    )
    assert (carlsen.standard_rating, carlsen.rapid_rating) == (2841, None)
    assert (
        carlsen.fide_id,
        carlsen.fide_standard_rating,
        carlsen.fide_rapid_rating,
        carlsen.fide_blitz_rating,
    ) == (1503014, 2823, 2803, 2860)
    anna = rows['NU0000001']
    assert (anna.club, anna.title, anna.gender, anna.year_of_birth) == (
        'FC ST.Pauli 1910 eV SAbt',
        'WFM',
        'F',
        2005,
    )
    doe = rows['NU0000002']
    assert (doe.club, doe.federation, doe.gender, doe.standard_rating) == (
        None,
        'GER',
        '',
        None,
    )


@pytest.mark.unit
def test_ecf_reads_the_rating_list(tmp_path: Path) -> None:
    from database.sqlite.national.ecf_database import EcfDatabase

    header = (
        'ECF_code,full_name,member_no,FIDE_no,gender,nation,original_standard,'
        'standard_original_category,revised_standard,standard_revised_category,'
        'original_rapid,rapid_original_category,revised_rapid,rapid_revised_category,'
        'original_blitz,blitz_original_category,revised_blitz,blitz_revised_category,'
        'original_standard_online,standard_online_original_category,'
        'revised_standard_online,standard_online_revised_category,'
        'original_rapid_online,rapid_online_original_category,revised_rapid_online,'
        'rapid_online_revised_category,original_blitz_online,'
        'blitz_online_original_category,revised_blitz_online,'
        'blitz_online_revised_category,club_code,club_name,title\r\n'
    )
    file = _write(
        tmp_path / 'ecf.csv',
        header
        + '100011B,"Arthurton, Robert J",19470,343434363,M,ENG,1614,A,1614,A,1571,K,'
        '1571,K,,,,,,,,,,,,,,,,,2SKG,Skegness,\r\n'
        + '109972D,"Dupré, Paul D",11992,405434,F,SCO,,,,,,,,,,,1300,C,,,,,,,,,,,,,'
        '1ABC,Wotton Hall,WIM\r\n',
    )
    rows = list(EcfDatabase()._read_players(file))
    robert, paul = rows
    assert (robert.national_id, robert.last_name, robert.first_name) == (
        '100011B',
        'Arthurton',
        'Robert J',
    )
    assert (robert.standard_rating, robert.rapid_rating, robert.blitz_rating) == (
        1614,
        1571,
        None,
    )
    assert (robert.fide_id, robert.club, robert.gender, robert.federation) == (
        343434363,
        'Skegness',
        'M',
        'ENG',
    )
    assert (paul.last_name, paul.blitz_rating, paul.title, paul.gender) == (
        'Dupré',
        1300,
        'WIM',
        'F',
    )
    assert paul.year_of_birth is None


@pytest.mark.unit
def test_fide_reads_the_players_out_of_the_xml_archive(tmp_path: Path) -> None:
    import zipfile

    from database.sqlite.fide.fide_database import FideDatabase

    xml = (
        '<playerslist>'
        '<player><fideid>1503014</fideid><name>Carlsen, Magnus</name>'
        '<country>NOR</country><sex>M</sex><title>GM</title><w_title></w_title>'
        '<o_title></o_title><rating>2823</rating><games>0</games><k>10</k>'
        '<rapid_rating>2803</rapid_rating><rapid_games>0</rapid_games>'
        '<rapid_k>10</rapid_k><blitz_rating>2860</blitz_rating>'
        '<blitz_games>0</blitz_games><blitz_k>10</blitz_k><birthday>1990</birthday>'
        '<flag></flag></player>'
        '<player><fideid>700070</fideid><name>Polgar, Judit</name>'
        '<country>HUN</country><sex>F</sex><title>WGM</title><w_title>WGM</w_title>'
        '<o_title>IA,FT</o_title><rating>2675</rating><games>0</games><k>10</k>'
        '<rapid_rating>0</rapid_rating><rapid_games>0</rapid_games><rapid_k>0</rapid_k>'
        '<blitz_rating>0</blitz_rating><blitz_games>0</blitz_games><blitz_k>0</blitz_k>'
        '<birthday>1976</birthday><flag>wi</flag></player>'
        '<player><fideid></fideid><name>Nobody</name></player>'
        '</playerslist>'
    )
    archive = tmp_path / 'players.zip'
    with zipfile.ZipFile(archive, 'w') as zf:
        zf.writestr('players_list_xml.xml', xml)
    database_file = tmp_path / 'fide.db'
    fide_database = FideDatabase()
    SQLiteDatabase(database_file, write=True)._create(fide_database._schema)
    with SQLiteDatabase(database_file, write=True) as database:
        assert fide_database._populate_from_source_file(archive, database)
        FideDatabase._create_indexes(database)
    with SQLiteDatabase(database_file) as database:
        database.execute(
            'SELECT fide_id, last_name, first_name, federation, gender, fide_title, '
            'fide_women_title, standard_rating, rapid_rating, blitz_rating, '
            'year_of_birth, k_standard, k_rapid, k_blitz, fide_arbiter_title '
            'FROM player ORDER BY fide_id'
        )
        rows = list(database.fetchall())
    assert [row['fide_id'] for row in rows] == [700070, 1503014]
    judit, magnus = rows
    assert (magnus['last_name'], magnus['first_name'], magnus['federation']) == (
        'Carlsen',
        'Magnus',
        'NOR',
    )
    assert (magnus['fide_title'], magnus['fide_women_title'], magnus['gender']) == (
        'GM',
        '',
        'M',
    )
    assert (
        magnus['standard_rating'],
        magnus['rapid_rating'],
        magnus['blitz_rating'],
        magnus['k_standard'],
    ) == (2823, 2803, 2860, 10)
    assert (judit['fide_title'], judit['fide_women_title']) == ('', 'WGM')
    assert (
        judit['fide_arbiter_title'],
        judit['year_of_birth'],
        judit['rapid_rating'],
    ) == (
        'IA',
        1976,
        0,
    )


@pytest.mark.unit
def test_chessa_reads_the_three_lists(tmp_path: Path) -> None:
    from database.sqlite.national.chessa_database import ChessaDatabase

    header = 'UNIQUE_NO,SURNAME,FIRSTNAME,BDATE,SEX,TITLE,RATING,FED\r\n'
    _write(
        tmp_path / 'ratings_rapid.csv',
        header + '182004314,Cawdery,Daniel John,1982/03/21,M,IM,2388,GAU\r\n',
    )
    _write(tmp_path / 'ratings_blitz.csv', header)
    file = _write(
        tmp_path / 'ratings_main.csv',
        header
        + '182004314,Cawdery,Daniel John,1982/03/21,M,IM,2305,GAU\r\n'
        + '100000016,A Nashir,Mohdkhairulnazr,1900/01/01,F,,1200,CSA\r\n',
    )
    daniel, nashir = list(ChessaDatabase()._read_players(file))
    assert (daniel.last_name, daniel.first_name, daniel.club, daniel.title) == (
        'Cawdery',
        'Daniel John',
        'GAU',
        'IM',
    )
    assert (daniel.standard_rating, daniel.rapid_rating, daniel.blitz_rating) == (
        2305,
        2388,
        None,
    )
    assert daniel.date_of_birth is not None and daniel.date_of_birth.isoformat() == (
        '1982-03-21'
    )
    assert (nashir.date_of_birth, nashir.year_of_birth, nashir.gender) == (
        None,
        None,
        'F',
    )


@pytest.mark.unit
def test_lok_reads_the_birthday_and_the_ids() -> None:
    from database.sqlite.national.lok_database import LokDatabase

    assert LokDatabase._birth('14.03.1985') == (1985, date(1985, 3, 14))
    assert LokDatabase._birth('00.00.2001') == (2001, None)
    assert LokDatabase._birth('') == (None, None)
    assert LokDatabase._birth('31.02.1990') == (1990, None)
    assert LokDatabase._cell_int(23760486.0) == 23760486
    assert LokDatabase._cell_int(0.0) is None
    assert LokDatabase._cell_int('') is None


@pytest.mark.unit
def test_ucf_splits_the_titles_off_the_names() -> None:
    from database.sqlite.national.ucf_database import UcfDatabase

    split = UcfDatabase._split_name
    assert split('Іванчук Василь мг/змс') == ('Іванчук', 'Василь', 'GM')
    assert split('Галинський Тимофій мм/мс') == ('Галинський', 'Тимофій', 'IM')
    assert split('Александров Валерій кмс') == ('Александров', 'Валерій', '')
    assert split('Жуков Володимир А. кмс') == ('Жуков', 'Володимир А.', '')
    assert split('Попович Андрій (мол)') == ('Попович', 'Андрій', '')
    assert split('Ель Кадер Марія') == ('Ель', 'Кадер Марія', '')


@pytest.mark.unit
def test_jcf_reads_the_rating_sheet_after_the_legend(tmp_path: Path) -> None:
    import openpyxl

    from database.sqlite.national.jcf_database import JcfDatabase

    workbook = openpyxl.Workbook()
    legend = workbook.active
    assert legend is not None
    legend.title = '項目説明'
    legend.append(['レーティングリストの項目について', None])
    legend.append(['ID', '日本チェス連盟会員ID'])
    sheet = workbook.create_sheet('2026-08-01')
    sheet.append(['2026-08-01', '国内レーティングリスト'])
    sheet.append(['ID', 'Name', 'Kanji', 'ST-K', 'ST', 'ST-G', 'RP-K', 'RP', 'RP-G'])
    sheet.append(['N08480036', 'Abdrakhmanov Andrei', 'X', 40, 1224, 0, 40, 0, 0])
    sheet.append(['N06752069', 'Zhu Shengyu', '朱 盛語', 40, 0, 0, 40, 1500, 3])
    file = tmp_path / 'list.xlsx'
    workbook.save(file)
    andrei, shengyu = list(JcfDatabase()._read_players(file))
    assert (andrei.national_id, andrei.last_name, andrei.first_name) == (
        'N08480036',
        'Abdrakhmanov',
        'Andrei',
    )
    assert (andrei.standard_rating, andrei.rapid_rating) == (1224, None)
    assert (shengyu.standard_rating, shengyu.rapid_rating) == (None, 1500)


@pytest.mark.unit
def test_nzcf_reads_the_two_header_rows(tmp_path: Path) -> None:
    import openpyxl

    from database.sqlite.national.nzcf_database import NzcfDatabase

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.append(
        ['Code', 'Name', None, 'Club', 'Year', None, None, 'Standard']
        + [None] * 6
        + ['Rapid']
        + [None] * 6
        + ['FIDE', None, None]
    )
    sheet.append(
        [
            None,
            None,
            None,
            None,
            None,
            None,
            None,
            'Rating',
            None,
            'T',
            'G',
            '∑G',
            'Active',
            None,
            'Rating',
            None,
            'T',
            'G',
            '∑G',
            'Active',
            None,
            'Code',
            'Fed',
            'Title',
        ]
    )
    sheet.append(
        [
            20069,
            'Abbott',
            'Isaac',
            '  ',
            2007,
            'J',
            None,
            '1545',
            '*',
            1,
            6,
            17,
            '2026-2',
            None,
            1411,
            ' ',
            ' ',
            ' ',
            29,
            '2026-1',
            None,
            4326350,
            'NZL',
            '   ',
        ]
    )
    sheet.append(
        [
            18452,
            'Abando',
            'Santino',
            'HB',
            1977,
            ' ',
            None,
            ' unr',
            ' ',
            '   ',
            '   ',
            '    ',
            '       ',
            None,
            1682,
            '*',
            ' ',
            ' ',
            6,
            '2022-4',
            None,
            4319591,
            'AUS',
            'FM',
        ]
    )
    file = tmp_path / 'list.xlsx'
    workbook.save(file)
    isaac, santino = list(NzcfDatabase()._read_players(file))
    assert (
        isaac.national_id,
        isaac.last_name,
        isaac.first_name,
        isaac.year_of_birth,
    ) == (
        '20069',
        'Abbott',
        'Isaac',
        2007,
    )
    assert (isaac.standard_rating, isaac.rapid_rating, isaac.fide_id, isaac.club) == (
        1545,
        1411,
        4326350,
        None,
    )
    assert (santino.standard_rating, santino.rapid_rating, santino.club) == (
        None,
        1682,
        'HB',
    )
    assert (santino.federation, santino.title) == ('AUS', 'FM')


@pytest.mark.unit
def test_dsu_reads_the_member_report(tmp_path: Path) -> None:
    from database.sqlite.national.dsu_database import DsuDatabase

    file = _write(
        tmp_path / 'members.csv',
        '""\r\n'
        '"Nr";"Født";"Alder";"Navn";"Klub";"Titel";"Rat.";"Hur.";"Lyn";"Fide";"Nr";"K"\r\n'
        '"103110502";"2004";"22";"Jonas Buhl Bjerre";"Skanderborg Skakklub";"GM";'
        '"2635";"2505";"2655";"2615";"1444948";"10"\r\n'
        '"102100848";"1980";"46";"Kunda Frederik Ollars";"Enkeltmedlem 9 HK";"";'
        '"0";"0";"0";"-";"0";""\r\n',
        'cp1252',
    )
    jonas, kunda = list(DsuDatabase()._read_players(file))
    assert (jonas.national_id, jonas.last_name, jonas.first_name) == (
        '103110502',
        'Bjerre',
        'Jonas Buhl',
    )
    assert (jonas.standard_rating, jonas.rapid_rating, jonas.blitz_rating) == (
        2635,
        2505,
        2655,
    )
    assert (jonas.fide_id, jonas.fide_standard_rating, jonas.title, jonas.club) == (
        1444948,
        2615,
        'GM',
        'Skanderborg Skakklub',
    )
    assert (kunda.standard_rating, kunda.fide_id, kunda.fide_standard_rating) == (
        None,
        None,
        None,
    )
    assert kunda.year_of_birth == 1980


@pytest.mark.unit
def test_percasi_reads_the_national_rating_sheet(tmp_path: Path) -> None:
    from database.sqlite.national.percasi_database import PercasiDatabase

    file = _write(
        tmp_path / 'drn.csv',
        ',,,,,,,,,\r\n'
        'DAFTAR RATING NASIONAL PER JULI 2026,,,,,,,,,\r\n'
        'ID.,FIDE-ID,PROVINSI,NAMA,L/P,Tgl.Lahir,Gelar,Standar,Cepat,Kilat\r\n'
        '10934,7102828,Aceh,ZULKHAIRI,,11/19/1975,FM,2169,2218,2268\r\n'
        '20982,,Aceh,Afifa Adelianda,P,,MN,1761,1760,1760\r\n'
        '"12345",,Jabar,"SUSANTO, DRS.",,,,1800,,\r\n',
    )
    zulkhairi, afifa, susanto = list(PercasiDatabase()._read_players(file))
    assert (zulkhairi.national_id, zulkhairi.last_name, zulkhairi.fide_id) == (
        '10934',
        'ZULKHAIRI',
        7102828,
    )
    assert zulkhairi.date_of_birth == date(1975, 11, 19)
    assert (zulkhairi.title, zulkhairi.gender, zulkhairi.club) == ('FM', 'M', 'Aceh')
    assert (
        zulkhairi.standard_rating,
        zulkhairi.rapid_rating,
        zulkhairi.blitz_rating,
    ) == (2169, 2218, 2268)
    assert (afifa.gender, afifa.title, afifa.year_of_birth) == ('F', '', None)
    assert (susanto.last_name, susanto.rapid_rating) == ('SUSANTO, DRS.', None)


@pytest.mark.unit
def test_mcf_reads_the_ids_and_gender() -> None:
    from database.sqlite.national.mcf_database import McfDatabase

    assert McfDatabase._cell_int('5707536') == 5707536
    assert McfDatabase._cell_int(0.0) is None
    assert McfDatabase._cell_int('') is None
